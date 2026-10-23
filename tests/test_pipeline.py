from datetime import date, datetime
from decimal import Decimal
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from urllib.error import HTTPError, URLError

from openpyxl import Workbook, load_workbook

from sales_report.cli import main
from sales_report.demo import serve
from sales_report.model import normalize
from sales_report.report import generate
from sales_report.sources import APIClient, SourceError, load_file
from sales_report.store import connect, import_rows, snapshot
from threading import Thread


ROW = dict(order_id='A1', order_date='2026-10-01', customer=' Demo Customer ', product='Router', quantity='2', unit_price='19.95', currency='usd')


class ValidationTests(unittest.TestCase):
    def test_exact_decimal_and_whitespace(self):
        order = normalize(ROW, 'USD')
        self.assertEqual(order.total, Decimal('39.90'))
        self.assertEqual(order.customer, 'Demo Customer')

    def test_invalid_values_are_flagged(self):
        for field, values in {
            'quantity': [0, -1, True, '2.5', '1000000'],
            'unit_price': ['NaN', 'Infinity', '-1', '1.005', 'USD 20', '1e100'],
            'order_date': ['01/02/2026', '2026-02-30', '2026-1-01'],
            'currency': ['EUR'], 'customer': ['', None, {'name': 'A'}, 'a\nb']
        }.items():
            for value in values:
                with self.subTest(field=field, value=value), self.assertRaises(ValueError):
                    normalize({**ROW, field: value}, 'USD')

    def test_excel_dates(self):
        for day in (date(2026, 10, 1), datetime(2026, 10, 1)):
            self.assertEqual(normalize({**ROW, 'order_date': day}, 'USD').order_date, '2026-10-01')


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.db = connect(self.root / 'orders.sqlite3')

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def test_idempotency_and_conflict_do_not_overwrite(self):
        first = import_rows(self.db, [ROW, ROW, {**ROW, 'quantity': '3'}], 'USD', 'test')
        run, orders = snapshot(self.db, first)
        stats = json.loads(run['stats'])
        self.assertEqual((stats['inserted'], stats['duplicates'], stats['rejected']), (1, 1, 1))
        self.assertEqual(orders[0].quantity, 2)
        second = import_rows(self.db, [ROW], 'USD', 'test')
        stats = json.loads(snapshot(self.db, second)[0]['stats'])
        self.assertEqual((stats['inserted'], stats['duplicates']), (0, 1))

    def test_wrong_database_currency_rolls_back(self):
        import_rows(self.db, [ROW], 'USD', 'test')
        with self.assertRaises(ValueError):
            import_rows(self.db, [{**ROW, 'order_id': 'B2', 'currency': 'EUR'}], 'EUR', 'test')
        self.assertEqual(self.db.execute('SELECT count(*) FROM orders').fetchone()[0], 1)
        self.assertEqual(self.db.execute('SELECT count(*) FROM runs').fetchone()[0], 1)

    def test_failed_import_rolls_back_every_insert(self):
        with self.assertRaises(AttributeError):
            import_rows(self.db, [ROW, None], 'USD', 'test')
        self.assertEqual(self.db.execute('SELECT count(*) FROM orders').fetchone()[0], 0)

    def test_reports_reconcile_and_preserve_text(self):
        row = {**ROW, 'customer': '=HYPERLINK("https://example.invalid")'}
        run = import_rows(self.db, [row, {**ROW, 'order_id': 'B2', 'order_date': 'ambiguous'}], 'USD', 'test')
        output, summary = generate(self.db, run, self.root / 'reports')
        self.assertEqual(summary['export']['total_sales'], '39.90')
        self.assertEqual(summary['export']['monthly_sales'], {'2026-10': '39.90'})
        self.assertEqual(summary['import']['received'], sum(summary['import'][k] for k in ('inserted', 'duplicates', 'rejected')))
        book = load_workbook(output / 'sales-report.xlsx')
        self.assertEqual(book.sheetnames, ['Summary', 'Clean orders', 'Import review'])
        self.assertEqual(book['Clean orders']['C2'].data_type, 's')
        self.assertEqual(book['Clean orders']['B2'].value.date(), date(2026, 10, 1))
        self.assertAlmostEqual(book['Clean orders']['H2'].value, 39.9)
        self.assertAlmostEqual(book['Summary']['B13'].value, 39.9)
        self.assertEqual(book['Import review'].max_row, 2)
        book.close()
        self.assertTrue((output / 'report.html').is_file())
        with self.assertRaises(ValueError):
            generate(self.db, run, self.root / 'reports')

    def test_aggregate_limit_rolls_back_and_keeps_reports_working(self):
        large = {**ROW, 'quantity': '1000', 'unit_price': '999999999.99'}
        first = import_rows(self.db, [large], 'USD', 'test')
        with self.assertRaisesRegex(ValueError, 'import rolled back'):
            import_rows(self.db, [{**ROW, 'order_id': 'B2'}], 'USD', 'test')
        self.assertEqual(self.db.execute('SELECT count(*) FROM orders').fetchone()[0], 1)
        self.assertEqual(self.db.execute('SELECT count(*) FROM runs').fetchone()[0], 1)
        _, summary = generate(self.db, first, self.root / 'reports')
        self.assertEqual(summary['export']['total_sales'], '999999999990.00')

    def test_rejected_control_character_exports_with_original_review(self):
        bad_id = 'A\x01B'
        run = import_rows(self.db, [ROW, {**ROW, 'order_id': bad_id}], 'USD', 'test')
        output, summary = generate(self.db, run, self.root / 'reports')
        self.assertEqual(summary['import']['rejected'], 1)
        book = load_workbook(output / 'sales-report.xlsx')
        self.assertEqual(book['Import review']['B2'].value, 'A\\u0001B')
        book.close()
        review = json.loads((output / 'review.json').read_text())
        self.assertEqual(review[0]['raw']['order_id'], bad_id)

    def test_empty_input_produces_zero_report(self):
        run = import_rows(self.db, [], 'USD', 'test')
        _, summary = generate(self.db, run, self.root / 'reports')
        self.assertEqual(summary['export']['total_sales'], '0')

    def test_sample_results_and_round_trip_excel(self):
        rows = load_file(Path('examples/orders.csv'))
        run = import_rows(self.db, rows, 'USD', 'orders.csv')
        output, summary = generate(self.db, run, self.root / 'reports')
        self.assertEqual((summary['import']['received'], summary['import']['inserted'], summary['import']['duplicates'], summary['import']['rejected']), (18, 12, 2, 4))
        self.assertEqual(summary['export']['total_sales'], '5645.20')
        # Input header schema is intentionally stable and explicit; exports use reader-friendly labels.
        with self.assertRaises(SourceError):
            load_file(output / 'sales-report.xlsx')

    def test_missing_headers_and_overlong_csv_rows_fail(self):
        for content in ('bad,headers\na,b\n', ','.join(ROW) + '\n' + ','.join(str(v) for v in ROW.values()) + ',extra\n'):
            path = self.root / 'bad.csv'
            path.write_text(content)
            with self.assertRaises(SourceError):
                load_file(path)

    def test_xlsx_import_accepts_typed_dates_and_numbers(self):
        path = self.root / 'orders.xlsx'
        book = Workbook()
        book.active.title = 'Orders'
        book.active.append(list(ROW))
        book.active.append(['XLSX-1', date(2026, 10, 1), 'Demo store', 'Router', 2, 19.95, 'USD'])
        book.save(path)
        book.close()
        rows = load_file(path, sheet='Orders')
        run = import_rows(self.db, rows, 'USD', 'orders.xlsx')
        self.assertEqual(snapshot(self.db, run)[1][0].total, Decimal('39.90'))
        with self.assertRaises(SourceError):
            load_file(path, sheet='Missing')

    def test_cli_api_failure_creates_no_database(self):
        path = self.root / 'not-created.sqlite3'
        with patch('sales_report.cli.APIClient.fetch', side_effect=SourceError('API request failed with HTTP 401')):
            with patch('sys.stderr'), patch('sys.stdout'):
                code = main(['--db', str(path), '--log', str(self.root/'run.log'),
                             'import', '--api', 'https://example.com/orders'])
        self.assertEqual(code, 1)
        self.assertFalse(path.exists())


class APITests(unittest.TestCase):
    def test_real_local_api_pagination_retries_and_rerun(self):
        server = serve(Path('examples/orders.csv'))
        worker = Thread(target=server.serve_forever, daemon=True)
        worker.start()
        try:
            client = APIClient(f'http://127.0.0.1:{server.server_port}/orders', sleep=lambda _: None)
            self.assertEqual(len(client.fetch()), 18)
            self.assertEqual(client.retries, 2)
            self.assertEqual(len(client.fetch()), 18)
        finally:
            server.shutdown()
            server.server_close()
            worker.join()

    def test_pagination_loop_is_not_a_partial_success(self):
        client = APIClient('https://example.com/orders')
        with patch.object(client, '_get', return_value={'items': [ROW], 'next_page': 1}), self.assertRaises(SourceError):
            client.fetch()

    def test_invalid_page_contract(self):
        for payload in ([], {'items': []}, {'items': [None], 'next_page': None}, {'items': [], 'next_page': True}):
            client = APIClient('https://example.com/orders')
            with self.subTest(payload=payload), patch.object(client, '_get', return_value=payload), self.assertRaises(SourceError):
                client.fetch()

    def test_unauthorized_fails_without_retry_or_secret_leak(self):
        client = APIClient('https://example.com/orders', sleep=lambda _: self.fail('must not retry 401'))
        error = HTTPError(client.url, 401, 'sensitive body', {}, None)
        with patch('sales_report.sources.build_opener') as opener:
            opener.return_value.open.side_effect = error
            with self.assertRaisesRegex(SourceError, 'HTTP 401') as raised:
                client.fetch()
            self.assertNotIn('sensitive', str(raised.exception))

    def test_retry_limit(self):
        delays = []
        client = APIClient('https://example.com/orders', attempts=3, sleep=delays.append)
        with patch('sales_report.sources.build_opener') as opener:
            opener.return_value.open.side_effect = URLError('secret detail')
            with self.assertRaisesRegex(SourceError, 'retry limit'):
                client.fetch()
        self.assertEqual(delays, [1, 2])

    def test_endpoint_rules_and_missing_token(self):
        for url in ('http://example.com/orders', 'https://user:password@example.com/orders', 'https://example.com/orders?token=secret'):
            with self.subTest(url=url), self.assertRaises(SourceError):
                APIClient(url)
        with patch.dict('os.environ', {}, clear=True), self.assertRaises(SourceError):
            APIClient('https://example.com/orders', token_env='MISSING_TOKEN')


if __name__ == '__main__':
    unittest.main()
