"""Input adapters. Network failures never return a partial successful import."""
import csv
from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
import json
import logging
import os
from pathlib import Path
import time
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener

from .model import FIELDS

LOG = logging.getLogger('sales_report')
MAX_ROWS = 100_000
MAX_BYTES = 10 * 1024 * 1024


class SourceError(ValueError):
    pass


def load_file(path: Path, sheet: str | None = None):
    if not path.is_file() or path.stat().st_size > 25 * 1024 * 1024:
        raise SourceError('Input must exist and be no larger than 25 MiB')
    rows = []
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            _headers(reader.fieldnames)
            for row in reader:
                if None in row:
                    raise SourceError('CSV row has more fields than its header')
                rows.append(row)
                _limit(rows)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet and sheet not in book.sheetnames:
                raise SourceError('Requested worksheet does not exist')
            selected = book[sheet] if sheet else book.worksheets[0]
            iterator = selected.iter_rows(values_only=True)
            headers = next(iterator, ())
            _headers(headers)
            for values in iterator:
                if all(value is None for value in values):
                    continue
                rows.append(dict(zip(headers, values)))
                _limit(rows)
        finally:
            book.close()
    else:
        raise SourceError('Supported input formats: .csv and .xlsx')
    return rows


def _headers(headers):
    if not headers or len(set(headers)) != len(headers) or not set(FIELDS).issubset(headers):
        raise SourceError('Input requires unique headers: ' + ', '.join(FIELDS))


def _limit(rows):
    if len(rows) > MAX_ROWS:
        raise SourceError(f'Input exceeds {MAX_ROWS} rows')


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):
        return None  # Never forward an Authorization header to a redirect destination.


class APIClient:
    """GET /orders?page=N -> {items: [...], next_page: N|null}."""
    def __init__(self, url, token_env=None, attempts=4, timeout=15, sleep=time.sleep):
        parsed = urlsplit(url)
        if parsed.scheme not in ('http', 'https') or not parsed.hostname:
            raise SourceError('API URL must use http or https')
        if parsed.username or parsed.password or parsed.query or parsed.fragment:
            raise SourceError('Use a URL without credentials, query parameters or fragments')
        if parsed.scheme != 'https' and parsed.hostname not in ('127.0.0.1', 'localhost', '::1'):
            raise SourceError('Remote API URLs must use HTTPS')
        if not 1 <= attempts <= 6 or not 0 < timeout <= 120:
            raise SourceError('Attempts must be 1–6 and timeout must be >0 and <=120 seconds')
        self.url, self.attempts, self.timeout, self.sleep = url, attempts, timeout, sleep
        self.headers = {'Accept': 'application/json', 'User-Agent': 'SalesReport/1.0'}
        if token_env:
            token = os.environ.get(token_env)
            if not token:
                raise SourceError('Configured API token environment variable is missing')
            if any(ord(char) < 33 or ord(char) > 126 for char in token):
                raise SourceError('API token must contain printable non-whitespace ASCII characters')
            self.headers['Authorization'] = 'Bearer ' + token
        self.retries = 0

    def _get(self, page):
        for attempt in range(self.attempts):
            delay = min(2 ** attempt, 30)
            try:
                request = Request(f'{self.url}?page={page}', headers=self.headers)
                with build_opener(NoRedirect()).open(request, timeout=self.timeout) as response:
                    data = response.read(MAX_BYTES + 1)
                if len(data) > MAX_BYTES:
                    raise SourceError('API page exceeds 10 MiB')
                try:
                    return json.loads(data)
                except (ValueError, UnicodeError) as exc:
                    raise SourceError('API returned invalid JSON') from exc
            except HTTPError as exc:
                status = exc.code
                retry_after = exc.headers.get('Retry-After', '')
                exc.close()
                if status not in (429, 500, 502, 503, 504):
                    raise SourceError(f'API request failed with HTTP {status}') from None
                if retry_after:
                    try:
                        delay = max(0, int(retry_after))
                    except ValueError:
                        try:
                            delay = max(0, (parsedate_to_datetime(retry_after) - datetime.now(timezone.utc)).total_seconds())
                        except (ValueError, TypeError, OverflowError):
                            pass
                    if delay > 60:
                        raise SourceError('Retry-After exceeds 60 seconds; rerun later') from None
            except (URLError, TimeoutError, OSError):
                status = 'network error'
            if attempt + 1 == self.attempts:
                raise SourceError('API retry limit exhausted; no rows imported')
            self.retries += 1
            LOG.warning('api_retry page=%s attempt=%s status=%s delay_seconds=%s', page, attempt + 1, status, delay)
            self.sleep(delay)

    def fetch(self):
        rows, seen, page = [], set(), 1
        for _ in range(1000):
            if page in seen:
                raise SourceError('API pagination loop detected')
            seen.add(page)
            payload = self._get(page)
            if not isinstance(payload, dict) or not isinstance(payload.get('items'), list) or 'next_page' not in payload:
                raise SourceError('API page must contain items list and next_page')
            if not all(isinstance(item, dict) for item in payload['items']):
                raise SourceError('Each API item must be an object')
            rows.extend(payload['items'])
            _limit(rows)
            page = payload['next_page']
            if page is None:
                return rows
            if type(page) is not int or page < 1:
                raise SourceError('next_page must be a positive integer or null')
        raise SourceError('API exceeded 1000 pages')
