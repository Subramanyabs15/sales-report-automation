"""Generate a workbook and a browser-readable summary from the same snapshot."""
from collections import defaultdict
from datetime import date
from decimal import Decimal
from html import escape
import json
from pathlib import Path
import shutil
import tempfile

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .store import snapshot


def generate(db, run_id, destination: Path):
    run, orders = snapshot(db, run_id)
    stats, issues = json.loads(run['stats']), json.loads(run['issues'])
    totals = defaultdict(lambda: Decimal('0'))
    for order in orders:
        totals[order.order_date[:7]] += order.total
    total = sum(totals.values(), Decimal('0'))
    if total > Decimal('999999999999.99'):
        raise ValueError('Report total exceeds supported Excel numeric precision')
    destination.mkdir(parents=True, exist_ok=True)
    final = destination / run_id
    if final.exists():
        raise ValueError('Report directory already exists; use another output directory to regenerate')
    temporary = Path(tempfile.mkdtemp(prefix='.report-', dir=destination))
    try:
        wb = Workbook()
        summary = wb.active
        summary.title = 'Summary'
        summary.append(['SALES REPORT', run['currency']])
        summary.append(['Run ID', run_id])
        summary.append(['Imported at (UTC)', run['created_at']])
        summary.append(['Import source', run['source']])
        summary.append(['Report scope', 'All stored orders at export; import counts refer to the selected run'])
        summary.append(['Metric', 'Value'])
        for key in ('received', 'inserted', 'duplicates', 'rejected', 'retries'):
            summary.append([key.replace('_', ' ').capitalize(), stats[key]])
        summary.append(['Orders in this export', len(orders)])
        summary.append(['Total sales', float(total)])
        summary['B13'].number_format = '#,##0.00'
        summary.append([])
        summary.append(['Month', 'Sales (' + run['currency'] + ')'])
        for month, value in sorted(totals.items()):
            summary.append([month, float(value)])
            summary.cell(summary.max_row, 2).number_format = '#,##0.00'
        clean = wb.create_sheet('Clean orders')
        clean.append(['Order ID', 'Order date', 'Customer', 'Product', 'Quantity', 'Unit price', 'Currency', 'Total'])
        for order in orders:
            clean.append([order.order_id, date.fromisoformat(order.order_date), order.customer, order.product,
                          order.quantity, float(Decimal(order.unit_price)), order.currency, float(order.total)])
        for row in range(2, clean.max_row + 1):
            clean.cell(row, 2).number_format = 'yyyy-mm-dd'
            for col in (6, 8):
                clean.cell(row, col).number_format = '#,##0.00'
        review = wb.create_sheet('Import review')
        review.append(['Input record', 'Order ID', 'Decision', 'Reason', 'Original record (JSON)'])
        for issue in issues:
            review.append([issue['record'], ILLEGAL_CHARACTERS_RE.sub(lambda match: f'\\u{ord(match.group()):04x}', issue['order_id']), issue['decision'], issue['reason'], json.dumps(issue['raw'], ensure_ascii=False, default=str)[:32000]])
        for sheet in wb:
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = 'A2'
            if sheet != summary:
                sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill('solid', fgColor='142D4E')
                cell.font = Font(name='Calibri', color='FFFFFF', bold=True, size=12)
            sheet.row_dimensions[1].height = 30
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = 'landscape'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_title_rows = '1:1'
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    cell.font = Font(name='Calibri', size=11, color='24364B')
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill('solid', fgColor='EFF4F8')
                sheet.row_dimensions[row[0].row].height = 26
            for col in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(col)].width = 22
        summary.column_dimensions['A'].width = 30
        summary.column_dimensions['B'].width = 80
        review.column_dimensions['D'].width = 65
        review.column_dimensions['E'].width = 85
        summary.freeze_panes = 'A7'
        summary.page_setup.fitToHeight = 1
        summary.row_dimensions[5].height = 38
        for row in range(2, review.max_row + 1):
            review.row_dimensions[row].height = 65
        wb.save(temporary / 'sales-report.xlsx')
        wb.close()
        result = {'run_id': run_id, 'created_at': run['created_at'], 'currency': run['currency'],
                  'import': stats, 'export': {'orders': len(orders), 'total_sales': str(total),
                  'monthly_sales': {month: str(value) for month, value in sorted(totals.items())}}}
        (temporary / 'summary.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
        (temporary / 'review.json').write_text(json.dumps(issues, indent=2, default=str), encoding='utf-8')
        _html(temporary / 'report.html', result, issues)
        temporary.rename(final)
    finally:
        if temporary.exists():
            shutil.rmtree(temporary)
    return final, result


def _html(path, result, issues):
    stats, export = result['import'], result['export']
    cards = ''.join(f'<article><span>{label}</span><strong>{value}</strong></article>' for label, value in [
        ('Records received', stats['received']), ('New orders', stats['inserted']),
        ('Duplicates skipped', stats['duplicates']), ('Needs review', stats['rejected'])])
    months = ''.join(f'<tr><td>{month}</td><td>{Decimal(value):,.2f}</td></tr>' for month, value in export['monthly_sales'].items())
    review = ''.join(f'<tr><td>{item["record"]}</td><td>{escape(item["order_id"])}</td><td>{escape(item["reason"])}</td></tr>' for item in issues)
    path.write_text('''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Sales automation report</title><style>
*{box-sizing:border-box}body{margin:0;background:#edf2f7;color:#172d48;font:16px system-ui,sans-serif}
main{max-width:1100px;margin:48px auto;padding:0 24px}header{background:#142d4e;color:white;padding:36px;border-radius:16px}
h1{font-size:36px;margin:10px 0}header p{color:#c6d7ed;max-width:750px;line-height:1.7}.eyebrow{letter-spacing:2px;font-size:12px}
.cards{display:grid;grid-template-columns:repeat(4,1fr);gap:16px;margin:24px 0}article,section{background:white;border:1px solid #dce4ed;border-radius:12px;padding:24px}
article span{display:block;color:#63748a;font-size:14px}strong{display:block;font-size:36px;margin-top:12px}section{margin-top:20px}
table{border-collapse:collapse;width:100%;font-size:14px}td,th{text-align:left;padding:14px 8px;border-bottom:1px solid #e6ebf1;overflow-wrap:anywhere}th{color:#63748a}
.note{color:#63748a;font-size:13px;line-height:1.7}.total{font-size:30px;font-weight:650}.badge{color:#6ee7c2}a{color:#16618e}footer{margin:24px 0;color:#63748a;font-size:13px}
@media(max-width:650px){.cards{grid-template-columns:repeat(2,1fr)}h1{font-size:26px}section{overflow:auto}}
</style><main><header><div class="eyebrow">PYTHON AUTOMATION / SALES OPERATIONS</div><h1>From raw orders to a reliable report.</h1><p>API and file imports with validation, persistent duplicate checks, retry handling, and a clear review trail.</p><div class="badge">Import completed</div></header>
''' + f'<div class="cards">{cards}</div><section><h2>Stored sales overview</h2><div class="total">{escape(result["currency"])} {Decimal(export["total_sales"]):,.2f}</div><p class="note">{export["orders"]} stored orders. Totals include all accepted records at export, not only this import. Taxes, refunds and currency conversion are outside this demo.</p><table><tr><th>Month</th><th>Sales ({escape(result["currency"])})</th></tr>{months}</table></section>'
        + f'<section><h2>Import review</h2><p class="note">{stats["retries"]} API retries. Conflicting records are flagged and never overwrite existing orders.</p><table><tr><th>Input record</th><th>Order ID</th><th>Reason</th></tr>{review}</table></section>'
        + f'<section><h2>Download outputs</h2><a href="sales-report.xlsx">Excel workbook</a> · <a href="summary.json">Summary JSON</a> · <a href="review.json">Review JSON</a></section><footer>Run {escape(result["run_id"])} · {escape(result["created_at"])}<br>Source and business rules are documented in the project README.</footer></main></html>', encoding='utf-8')
